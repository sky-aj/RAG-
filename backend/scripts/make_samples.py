r"""生成 DOCX / PDF 样例文件。

这两个加载器没法手写样例（二进制格式），所以用代码生成，
样例文件本身就是「可以查看到的中间产物」，方便对着验证解析结果。

用法：.venvScriptspython.exe scriptsmake_samples.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, OSError):
        pass

SAMPLES = Path(__file__).resolve().parent.parent / "data" / "samples"

# 用系统里能拿到的中文字体，拿不到就退到 Helvetica（英文照样能验证结构）
CJK_CANDIDATES = [
    (Path(r"C:\Windows\Fonts") / "simhei.ttf", "SimHei"),
    (Path(r"C:\Windows\Fonts") / "simkai.ttf", "KaiTi"),
    (Path(r"C:\Windows\Fonts") / "Deng.ttf", "DengXian"),
]


def build_docx() -> Path:
    from docx import Document
    from docx.shared import Pt

    document = Document()

    document.add_heading("运维 Runbook", level=1)
    document.add_paragraph("本手册覆盖发布、灰度、回滚与容量评估四类操作，供值班同学直接照着执行。")

    document.add_heading("一、发布流程", level=2)
    document.add_paragraph("发布前确认三件事：灰度开关已就位、数据库迁移脚本已评审、回滚包已构建。")
    document.add_paragraph("任何一项缺失都不应进入发布窗口。")

    document.add_heading("二、容量评估", level=2)
    document.add_paragraph("容量按峰值 QPS 的 1.4 倍预留，单实例上限如下：")
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    header = ["节点类型", "QPS 上限", "说明"]
    for index, text in enumerate(header):
        table.cell(0, index).text = text
    rows = [
        ["API 节点", "1200", "8 核 16G"],
        ["检索节点", "800", "受向量库限制"],
        ["索引节点", "200", "批量任务用"],
    ]
    for row_index, row in enumerate(rows, start=1):
        for cell_index, text in enumerate(row):
            table.cell(row_index, cell_index).text = text

    document.add_heading("三、回滚", level=2)
    document.add_paragraph("索引重建期间建议先将流量切到只读副本，重建完成后再切换回主节点。")
    document.add_paragraph("回滚包保留最近三个版本，超过三个版本的包会被自动清理。")

    paragraph = document.add_paragraph()
    run = paragraph.add_run("注意：回滚后必须核对指标看板，确认错误率回到基线再解除封禁。")
    run.bold = True
    run.font.size = Pt(11)

    path = SAMPLES / "运维 Runbook.docx"
    document.save(str(path))
    return path


def build_pdf() -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors

    font_name = "Helvetica"
    registered: list[str] = []
    for font_path, name in CJK_CANDIDATES:
        if not font_path.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont(name, str(font_path)))
            font_name = name
            registered.append(name)
            break
        except Exception as error:  # noqa: BLE001 - 字体注册失败就换下一个
            print(f"  警告：字体 {font_path.name} 注册失败（{error}）")
            continue

    if not registered:
        # 早期这里静默退回 Helvetica，生成的 PDF 汉字全变 ■，
        # 排查时误以为解析器有问题，所以现在必须显式告警
        print("  警告：没有可用中文字体，样例 PDF 里的中文会变成方块")
        print("        这属于样例生成问题，不是解析器问题")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DocTitle", parent=styles["Title"], fontName=font_name, fontSize=20)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName=font_name, fontSize=14)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16)
    cell = ParagraphStyle("Cell", parent=body, fontSize=9, leading=12)

    path = SAMPLES / "API 参考手册.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, title="API 参考手册")
    story = [
        Paragraph("API 参考手册", title_style),
        Spacer(1, 14),
        Paragraph("本文档描述开放平台的检索接口与限流约定。", body),
        Spacer(1, 14),
        Paragraph("1. 检索接口", h2),
        Spacer(1, 8),
        Paragraph("GET /v1/search 支持 top_k 与 score_threshold 两个参数；"
                  "score_threshold 低于 0.35 的片段默认被丢弃，可通过查询参数覆盖。", body),
        Spacer(1, 8),
        Paragraph("返回体里的 score 是余弦相似度，越接近 1 越相关。", body),
        Spacer(1, 16),
        Paragraph("2. 限流与配额", h2),
        Spacer(1, 8),
        Paragraph("对外接口按租户维度限流，默认额度为每分钟 600 次请求、并发 20。"
                  "额度以滑动窗口统计，超出后返回 429 并附 Retry-After 响应头。", body),
        Spacer(1, 12),
        Table(
            [[Paragraph("场景", cell), Paragraph("限额", cell), Paragraph("说明", cell)],
             [Paragraph("默认", cell), Paragraph("600 次/分钟", cell), Paragraph("所有租户", cell)],
             [Paragraph("突发", cell), Paragraph("900 次/分钟", cell), Paragraph("持续不超过 10 秒", cell)],
             [Paragraph("企业版", cell), Paragraph("3000 次/分钟", cell), Paragraph("需申请", cell)]],
            colWidths=[90, 120, 200],
            style=TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]),
        ),
        Spacer(1, 16),
        Paragraph("3. 幂等", h2),
        Spacer(1, 8),
        Paragraph("所有写操作接口要求 Idempotency-Key 请求头，重复的 key 在 24 小时内"
                  "返回同一结果，避免网络重试造成重复扣费。", body),
        PageBreak(),
        Paragraph("4. 错误码", h2),
        Spacer(1, 8),
        Paragraph("400 参数错误；401 鉴权失败；429 触发限流；500 服务端异常。", body),
        Spacer(1, 8),
        Paragraph("429 响应一定会带 Retry-After，单位是秒。", body),
    ]
    doc.build(story)
    return path


def main() -> int:
    print("=" * 62)
    print("生成 DOCX / PDF 样例")
    print("=" * 62)
    docx_path = build_docx()
    print(f"  DOCX: {docx_path.name}  {docx_path.stat().st_size / 1024:.1f} KB")
    pdf_path = build_pdf()
    print(f"  PDF : {pdf_path.name}  {pdf_path.stat().st_size / 1024:.1f} KB")
    print("完成。可以跑 scripts\\smoke_step2.py 对照解析结果。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
